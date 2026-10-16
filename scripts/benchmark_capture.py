#!/usr/bin/env python3
"""Credit-careful benchmark capture for the V2 waterfall (LYCERT_WATERFALL_ENRICHMENT_V2 Part 8.2).

ONE network pass: runs each provider with a key independently against the ground-truth sample,
writes every raw returned field + the expected ground-truth values to a gitignored JSON capture, then
prints (a) the Part 8.2 per-provider table and (b) a per-subscriber Clay-vs-ours field comparison.
Because the raw is persisted, the analysis can be recomputed offline at zero credit cost.

  python scripts/benchmark_capture.py "C:/path/Ground truth sample.xlsx"          # real (uses credits)
  python scripts/benchmark_capture.py "C:/path/Ground truth sample.xlsx" --dry    # no network, stubs

NOTE: 5 rows is an indicative spot-check, NOT a basis to reorder the live cascade. Record qualitative
observations only; do not touch enrichment_providers.default_position from a sample this small.
"""

import asyncio
import json
import os
import sys
from pathlib import Path

_REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(_REPO_ROOT / "apps" / "api"))

CAPTURE_PATH = _REPO_ROOT / "bench_capture.json"  # gitignored


def _load_dotenv() -> None:
    for fname in (".env.local", ".env"):
        path = _REPO_ROOT / fname
        if not path.exists():
            continue
        for raw in path.read_text(encoding="utf-8").splitlines():
            line = raw.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            if line.startswith("export "):
                line = line[len("export "):]
            k, _, v = line.partition("=")
            k, v = k.strip(), v.strip().strip('"').strip("'")
            if k and k not in os.environ:
                os.environ[k] = v


# Ground-truth (the design partner) column name → our internal key.
_COLUMN_ALIASES = {
    "contact name": "name",
    "linkedin profile final": "expected_linkedin",
    "li title": "expected_title",
    "li job company name": "expected_employer",
    "company domain name": "expected_domain",
    "li education name": "expected_school",
    "li education date range": "expected_edu_range",
    "li job start date": "expected_job_start",
    "li job end date": "expected_job_end",
}

# Returned canonical field → the expected ground-truth key it should be compared against.
_COMPARE = [
    ("LinkedIn", "linkedin_profile_url", "expected_linkedin"),
    ("Title", "li_job_title", "expected_title"),
    ("Employer", "li_current_employer_name", "expected_employer"),
    ("Emp domain", "li_current_employer_domain", "expected_domain"),
    ("School", "li_education_institution", "expected_school"),
    ("Edu end", "li_education_end_date", "expected_edu_range"),
    ("Job start", "li_job_start_date", "expected_job_start"),
    ("Job end", "li_job_end_date", "expected_job_end"),
]


def _load_sample(path: Path) -> list[dict]:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    headers = [str(c).strip().lower() if c is not None else "" for c in next(rows_iter)]
    out = []
    for r in rows_iter:
        row = {}
        for i, h in enumerate(headers):
            if not h or i >= len(r):
                continue
            val = r[i]
            row[h] = "" if val is None else str(val).strip()
            if h in _COLUMN_ALIASES and _COLUMN_ALIASES[h] not in row:
                row[_COLUMN_ALIASES[h]] = row[h]
        if row.get("email"):
            out.append(row)
    return out


def _providers():
    from services.enrichment.providers.datagma_client import DatagmaClient
    from services.enrichment.providers.hunter_client import HunterClient
    from services.enrichment.providers.pdl_client import PDLProvider
    from services.enrichment.providers.reversecontact_client import ReverseContactClient
    return [
        ("hunter", HunterClient, "LYCERT_HUNTER_API_KEY", 0.004),
        ("reversecontact", ReverseContactClient, "LYCERT_REVERSECONTACT_API_KEY", 0.030),
        ("datagma", DatagmaClient, "LYCERT_DATAGMA_API_KEY", 0.030),
        ("pdl", PDLProvider, "LYCERT_PDL_API_KEY", 0.280),
    ]


async def _run_real(sample: list[dict], only: set[str] | None = None) -> dict:
    captured: dict = {}
    for name, cls, env_var, cost in _providers():
        if only and name not in only:
            continue  # --only filter: skip providers not requested (avoids re-spending on PDL etc.)
        if not os.environ.get(env_var):
            print(f"  SKIP {name}: {env_var} not set")
            continue
        print(f"  Running {name} against {len(sample)} subscribers...")
        client = cls()
        per_sub = []
        for i, row in enumerate(sample):
            res = await client.enrich(row["email"], row.get("name"), f"bench-{i}")
            per_sub.append({
                "email": row["email"],
                "outcome": getattr(res.outcome, "value", str(res.outcome)),
                "charged": bool(res.charged),
                "cost_usd": float(res.cost_usd or 0),
                "full_name": res.full_name,
                "error": res.error_message,
                "fields": res.fields or {},
            })
        captured[name] = {"cost_planning": cost, "results": per_sub}
    return captured


def _run_dry(sample: list[dict]) -> dict:
    """No-network stub so the loader + comparison formatting can be verified at zero cost."""
    captured = {}
    for name, _cls, _env, cost in _providers():
        per_sub = []
        for row in sample:
            per_sub.append({
                "email": row["email"], "outcome": "partial", "charged": True, "cost_usd": cost,
                "full_name": row.get("name"),
                "error": None,
                "fields": {"linkedin_profile_url": "https://linkedin.com/in/stub",
                           "li_current_employer_name": "StubCo", "li_job_title": "Stub Title"},
            })
        captured[name] = {"cost_planning": cost, "results": per_sub}
    return captured


def _summarize(captured: dict, n: int) -> None:
    print("\n| Provider | Hit rate | Edu cov | Dates cov | $/hit |")
    print("|---|---|---|---|---|")
    for name, blob in captured.items():
        rs = blob["results"]
        hits = [r for r in rs if r["outcome"] in ("complete", "partial")]
        edu = sum(1 for r in hits if r["fields"].get("li_education_end_date"))
        dates = sum(1 for r in hits if r["fields"].get("li_job_start_date"))
        spend = sum(r["cost_usd"] for r in rs)
        cph = (spend / len(hits)) if hits else blob["cost_planning"]
        ec = f"{edu/len(hits):.0%}" if hits else "—"
        dc = f"{dates/len(hits):.0%}" if hits else "—"
        print(f"| {name} | {len(hits)/n:.0%} ({len(hits)}/{n}) | {ec} | {dc} | ${cph:.4f} |")


def _val(x) -> str:
    return (x if isinstance(x, str) else str(x)) if x else ""


def _compare(captured: dict, sample: list[dict]) -> None:
    by_email = {r["email"]: r for r in sample}
    for email, exp in by_email.items():
        print(f"\n=== {email}  ({exp.get('name','?')}) ===")
        for label, ret_key, exp_key in _COMPARE:
            expected = _val(exp.get(exp_key))
            cells = []
            for pname, blob in captured.items():
                row = next((r for r in blob["results"] if r["email"] == email), None)
                got = _val(row["fields"].get(ret_key)) if row else ""
                if got:
                    cells.append(f"{pname}={got}")
            ours = "; ".join(cells) if cells else "—"
            if not expected and not cells:
                continue
            bucket = ("both" if expected and cells else
                      "CLAY-ONLY (ours missing)" if expected and not cells else
                      "OURS-ONLY (not in Clay)" if cells and not expected else "—")
            print(f"  {label:11} | clay: {expected or '—'}")
            print(f"  {'':11} | ours: {ours}    [{bucket}]")


async def main() -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
    except Exception:
        pass
    _load_dotenv()
    dry = "--dry" in sys.argv
    # --only datagma  (or --only=hunter,datagma) → run just those providers (don't re-spend on others).
    only: set[str] | None = None
    rest: list[str] = []
    argv = sys.argv[1:]
    i = 0
    while i < len(argv):
        a = argv[i]
        if a == "--dry":
            pass
        elif a == "--only" and i + 1 < len(argv):
            only = {p.strip() for p in argv[i + 1].split(",") if p.strip()}
            i += 1
        elif a.startswith("--only="):
            only = {p.strip() for p in a.split("=", 1)[1].split(",") if p.strip()}
        else:
            rest.append(a)
        i += 1
    args = rest
    if not args:
        print("Usage: python scripts/benchmark_capture.py <ground_truth.xlsx> [--dry] [--only datagma]")
        return 1
    path = Path(args[0])
    if not path.exists():
        print(f"Sample not found: {path}")
        return 1
    sample = _load_sample(path)
    print(f"Loaded {len(sample)} subscribers from {path.name}{' (DRY — no network)' if dry else ''}\n")

    if only:
        print(f"--only filter active: {sorted(only)}\n")
    captured = _run_dry(sample) if dry else await _run_real(sample, only)
    # default=str: some providers return date/datetime objects in fields — keep the capture robust
    # so a single run is never lost to a serialization error (the whole point of capturing raw).
    CAPTURE_PATH.write_text(
        json.dumps({"sample": sample, "captured": captured}, indent=2, default=str), encoding="utf-8")
    print(f"\nRaw capture written to {CAPTURE_PATH.name} (gitignored).")

    _summarize(captured, len(sample))
    _compare(captured, sample)
    print("\n⚠ 5-row spot-check — qualitative only. Do NOT reorder enrichment_providers.default_position.")
    return 0


if __name__ == "__main__":
    raise SystemExit(asyncio.run(main()))
