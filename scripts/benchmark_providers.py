#!/usr/bin/env python3
"""Benchmark enrichment providers against the design-partner ground-truth sample.

LYCERT_WATERFALL_ENRICHMENT_V2.md Part 8.2 — THE BENCHMARK. The required first action after
build, BEFORE signing any provider contract.

What it does:
  - Loads a sample of subscribers (emails + known-correct LinkedIn / education / employment).
  - Runs each provider's free tier INDEPENDENTLY (not cascaded) against the whole sample.
  - Prints a per-provider table: hit rate, education coverage, dates coverage, wrong-person
    count, effective cost/hit.
  - Applies the decision rule: rank by (hit_rate × completeness) ÷ cost; flag any provider
    under 15% measured hit rate for dropping.
  - You then record the results + final order in DECISIONS.md and update
    enrichment_providers.default_position accordingly.

Usage:
  python scripts/benchmark_providers.py [path/to/sample.csv|.xlsx]

Sample columns (case-insensitive; only `email` is required):
  email, name, expected_linkedin, expected_school, expected_employer

Requires the provider API keys to be set in the environment (free-tier keys are enough:
ZeroBounce 100/mo, Hunter 25/mo, ReverseContact trial, Datagma trial, PDL 100/mo).
FullEnrich is async (webhook delivery) and is NOT part of this inline benchmark.
"""

import asyncio
import csv
import os
import sys
from pathlib import Path

# Make the FastAPI app's modules importable.
_REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(_REPO_ROOT / "apps" / "api"))


def _load_dotenv() -> None:
    """Populate os.environ from a gitignored .env.local (or .env) at the repo root, so provider keys
    pulled from Vercel via `vercel env pull .env.local` are picked up without a python-dotenv
    dependency. A value already present in the real shell environment always wins (never overridden).
    Minimal KEY=VALUE parser: skips blanks/comments, strips surrounding quotes, ignores `export `.
    """
    for fname in (".env.local", ".env"):
        path = _REPO_ROOT / fname
        if not path.exists():
            continue
        for raw in path.read_text(encoding="utf-8").splitlines():
            line = raw.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            if line.startswith("export "):
                line = line[len("export "):]
            key, _, val = line.partition("=")
            key = key.strip()
            val = val.strip().strip('"').strip("'")
            if key and key not in os.environ:
                os.environ[key] = val

from services.enrichment.merge import names_overlap  # noqa: E402
from services.enrichment.providers.base import ProviderOutcome  # noqa: E402
from services.enrichment.providers.datagma_client import DatagmaClient  # noqa: E402
from services.enrichment.providers.hunter_client import HunterClient  # noqa: E402
from services.enrichment.providers.pdl_client import PDLProvider  # noqa: E402
from services.enrichment.providers.reversecontact_client import ReverseContactClient  # noqa: E402

# Provider key → (class, env var, planning cost/hit). FullEnrich excluded (async/webhook).
_PROVIDERS = [
    ("hunter", HunterClient, "LYCERT_HUNTER_API_KEY", 0.004),
    ("reversecontact", ReverseContactClient, "LYCERT_REVERSECONTACT_API_KEY", 0.030),
    ("datagma", DatagmaClient, "LYCERT_DATAGMA_API_KEY", 0.030),
    ("pdl", PDLProvider, "LYCERT_PDL_API_KEY", 0.280),
]

_MIN_HIT_RATE = 0.15  # drop any provider under 15% measured hit rate (Part 8.2 decision rule)


def _load_sample(path: Path) -> list[dict]:
    rows: list[dict] = []
    if path.suffix.lower() in (".xlsx", ".xls"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        headers = [str(c.value).strip().lower() if c.value else "" for c in next(ws.iter_rows(max_row=1))]
        for r in ws.iter_rows(min_row=2, values_only=True):
            rows.append({headers[i]: r[i] for i in range(len(headers)) if i < len(r)})
    else:
        with open(path, newline="", encoding="utf-8") as f:
            for r in csv.DictReader(f):
                rows.append({(k or "").strip().lower(): v for k, v in r.items()})
    return [r for r in rows if r.get("email")]


async def _bench_provider(name: str, cls, cost: float, sample: list[dict]) -> dict:
    client = cls()
    hits = edu = dates = wrong = charged_calls = 0
    total_cost = 0.0
    for i, row in enumerate(sample):
        result = await client.enrich(row["email"], row.get("name"), f"bench-{i}")
        if result.charged:
            charged_calls += 1
            total_cost += result.cost_usd
        if result.outcome in (ProviderOutcome.COMPLETE, ProviderOutcome.PARTIAL):
            hits += 1
            f = result.fields
            if f.get("li_education_end_date"):
                edu += 1
            if f.get("li_job_start_date"):
                dates += 1
            expected_name = row.get("name")
            if expected_name and result.full_name and not names_overlap(expected_name, result.full_name):
                wrong += 1
    n = len(sample)
    hit_rate = hits / n if n else 0.0
    completeness = ((edu / hits) + (dates / hits)) / 2 if hits else 0.0
    cost_per_hit = (total_cost / hits) if hits else cost
    score = (hit_rate * completeness) / cost_per_hit if cost_per_hit else 0.0
    return {
        "provider": name, "hits": hits, "n": n, "hit_rate": hit_rate,
        "edu_cov": (edu / hits) if hits else 0.0, "dates_cov": (dates / hits) if hits else 0.0,
        "wrong": wrong, "cost_per_hit": cost_per_hit, "score": score,
    }


async def main() -> int:
    # Windows consoles default to cp1252, which can't encode the table's em-dash / ⚠ glyphs and
    # crashes the run mid-print. Force UTF-8 stdout where supported.
    try:
        sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
    except Exception:
        pass
    _load_dotenv()  # pick up keys pulled into .env.local from Vercel
    default_sample = _REPO_ROOT / "docs" / "reference" / "sample_accounts.xlsx"
    sample_path = Path(sys.argv[1]) if len(sys.argv) > 1 else default_sample
    if not sample_path.exists():
        print(f"Sample file not found: {sample_path}")
        return 1
    if sample_path.resolve() == default_sample.resolve():
        # The CEDS scoring-regression fixture holds SYNTHETIC placeholder emails (johnsmith@gmail.com …)
        # with no real person behind them, so every provider scores 0% and the benchmark is meaningless.
        # Part 8.2 requires the real design-partner Clay-enriched ground-truth Excel (real emails + known-correct
        # LinkedIn/education/employment). Pass its path explicitly.
        print("⚠ WARNING: using the CEDS regression fixture (sample_accounts.xlsx) — SYNTHETIC emails.")
        print("  Provider hits will be 0% and meaningless. Pass the real design-partner ground-truth file:")
        print("  python scripts/benchmark_providers.py path/to/partner_ground_truth.xlsx\n")
    sample = _load_sample(sample_path)
    print(f"Loaded {len(sample)} subscribers from {sample_path.name}\n")

    results = []
    for name, cls, env_var, cost in _PROVIDERS:
        if not os.environ.get(env_var):
            print(f"  SKIP {name}: {env_var} not set")
            continue
        print(f"  Running {name} against {len(sample)} subscribers...")
        results.append(await _bench_provider(name, cls, cost, sample))

    if not results:
        print("\nNo providers ran — set free-tier API keys and retry.")
        return 1

    print("\n| Provider | Hit rate | Edu cov | Dates cov | Wrong-person | $/hit | Score |")
    print("|---|---|---|---|---|---|---|")
    for r in sorted(results, key=lambda x: x["score"], reverse=True):
        print(f"| {r['provider']} | {r['hit_rate']:.0%} ({r['hits']}/{r['n']}) | "
              f"{r['edu_cov']:.0%} | {r['dates_cov']:.0%} | {r['wrong']} | "
              f"${r['cost_per_hit']:.4f} | {r['score']:.2f} |")

    print("\nDecision rule — recommended default_position order (best score first):")
    for pos, r in enumerate(sorted(results, key=lambda x: x["score"], reverse=True), start=1):
        flag = "  ⚠ DROP (< 15% hit rate)" if r["hit_rate"] < _MIN_HIT_RATE else ""
        print(f"  {pos}. {r['provider']}{flag}")
    print("\nRecord these results and the final order in DECISIONS.md, then update "
          "enrichment_providers.default_position.")
    return 0


if __name__ == "__main__":
    raise SystemExit(asyncio.run(main()))
