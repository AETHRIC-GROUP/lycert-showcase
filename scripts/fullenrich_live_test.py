#!/usr/bin/env python3
"""Live FullEnrich verification + Clay comparison (poll-only, no webhook needed).

Submits the ground-truth sample to FullEnrich's reverse-email bulk endpoint and POLLS for results
(the GET payload is the same data[] shape as the webhook), then prints a masked Clay-vs-ours
comparison. Writes the raw result to a gitignored capture so the real response STRUCTURE can be
reconciled with the TS parser offline — re-readable at zero credit cost.

  python scripts/fullenrich_live_test.py "C:/path/Ground truth sample.xlsx"

Uses LYCERT_FULLENRICH_API_KEY from .env.local. Charges 1 credit per successful match (≤5 here).
NO webhook_url is sent (poll-only) to avoid 401-spam before LYCERT_FULLENRICH_WEBHOOK_SECRET is set.
"""
import json
import os
import sys
import time
from pathlib import Path

import httpx

_REPO_ROOT = Path(__file__).resolve().parents[1]
CAPTURE = _REPO_ROOT / "bench_capture_fullenrich.json"  # gitignored (bench_capture*.json)
BASE = "https://app.fullenrich.com/api/v2"


def _load_dotenv() -> None:
    for fname in (".env.local", ".env"):
        p = _REPO_ROOT / fname
        if not p.exists():
            continue
        for raw in p.read_text(encoding="utf-8").splitlines():
            line = raw.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            k, _, v = line.partition("=")
            if k.strip() and k.strip() not in os.environ:
                os.environ[k.strip()] = v.strip().strip('"').strip("'")


def _load_sample(path: Path) -> list[dict]:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    headers = [str(c).strip().lower() if c is not None else "" for c in next(it)]
    rows = []
    for r in it:
        row = {headers[i]: ("" if r[i] is None else str(r[i]).strip()) for i in range(len(headers)) if i < len(r)}
        if row.get("email"):
            rows.append(row)
    return rows


def mask(e: str) -> str:
    l, _, d = e.partition("@")
    return (l[:1] + "***@" + d) if l else e


def extract(profile: dict) -> dict:
    """Mirror of TS mapFullenrichProfile — for the human comparison only."""
    if not profile:
        return {}
    sp = (profile.get("social_profiles") or {}).get("professional_network") or {}
    emp = profile.get("employment") or {}
    cur = emp.get("current") or (next((e for e in (emp.get("all") or []) if e.get("is_current")), None)) or (emp.get("all") or [{}])[0] if emp.get("all") else (emp.get("current") or {})
    cur = cur or {}
    comp = cur.get("company") or {}
    edus = profile.get("educations") or []
    edu = max(edus, key=lambda e: (e.get("end_at") or ""), default={}) if edus else {}
    loc = profile.get("location") or {}
    return {
        "linkedin": sp.get("url"),
        "title": cur.get("title"),
        "employer": comp.get("name"),
        "domain": comp.get("domain"),
        "school": edu.get("school_name"),
        "edu_end": (edu.get("end_at") or "")[:4] or None,
        "job_start": cur.get("start_at"),
        "country": loc.get("country"),
    }


def main() -> int:
    _load_dotenv()
    key = os.environ.get("LYCERT_FULLENRICH_API_KEY")
    if not key:
        print("LYCERT_FULLENRICH_API_KEY not set"); return 1
    if len(sys.argv) < 2:
        print('Usage: python scripts/fullenrich_live_test.py "<ground truth>.xlsx"'); return 1
    sample = _load_sample(Path(sys.argv[1]))
    print(f"Loaded {len(sample)} subscribers\n")

    headers = {"Authorization": f"Bearer {key}", "Content-Type": "application/json"}
    body = {
        "name": "lycert_live_test",
        "data": [{"email": r["email"], "custom": {"lycert_subscriber_id": r["email"]}} for r in sample],
    }
    sub = httpx.post(f"{BASE}/contact/reverse/email/bulk", headers=headers, json=body, timeout=60)
    print("submit:", sub.status_code, sub.text[:300])
    if sub.status_code != 200:
        return 1
    enrichment_id = sub.json().get("enrichment_id")
    if not enrichment_id:
        print("no enrichment_id"); return 1
    print("enrichment_id:", enrichment_id)

    result = None
    for attempt in range(20):  # ~160s max
        time.sleep(8)
        g = httpx.get(f"{BASE}/contact/reverse/email/bulk/{enrichment_id}", headers=headers, timeout=60)
        if g.status_code != 200:
            print(f"  poll {attempt}: HTTP {g.status_code}"); continue
        result = g.json()
        status = result.get("status")
        print(f"  poll {attempt}: status={status} credits={(result.get('cost') or {}).get('credits')}")
        if status == "FINISHED":
            break

    if not result:
        print("no result"); return 1
    CAPTURE.write_text(json.dumps(result, indent=2, default=str), encoding="utf-8")
    print(f"\nRaw result captured to {CAPTURE.name} (gitignored).\n")

    by_email = {r["email"]: r for r in sample}
    for item in result.get("data") or []:
        email = (item.get("custom") or {}).get("lycert_subscriber_id") or (item.get("input") or {}).get("email") or "?"
        exp = by_email.get(email, {})
        ours = extract(item.get("profile") or {})
        print(f"=== {mask(email)} ===  (profile present: {bool(item.get('profile'))})")
        rows = [
            ("LinkedIn", exp.get("linkedin profile final"), ours.get("linkedin")),
            ("Title", exp.get("li title"), ours.get("title")),
            ("Employer", exp.get("li job company name"), ours.get("employer")),
            ("Emp domain", exp.get("company domain name"), ours.get("domain")),
            ("School", exp.get("li education name"), ours.get("school")),
            ("Edu end", None, ours.get("edu_end")),
            ("Job start", None, ours.get("job_start")),
        ]
        for label, clay, got in rows:
            if not clay and not got:
                continue
            print(f"  {label:11} clay: {clay or '—'}")
            print(f"  {'':11} ours: {got or '—'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
